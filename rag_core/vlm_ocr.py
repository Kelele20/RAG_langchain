# -*- coding: utf-8 -*-
"""
多模态 OCR + VLM 文档解析模块
- 支持 VLM (Vision-Language Model) 视觉大模型对图片、图表、流程图、扫描件进行高保真结构化 Markdown 提取。
- 支持 PDF、Word、Excel、CSV、JSON、TXT、Markdown 等多格式全自动化加载与文本规范化。
"""

import os
import io
import base64
import json
import pandas as pd
import pypdf
import docx
from typing import List, Dict, Any, Optional, Tuple
from langchain_core.documents import Document


def encode_image_to_base64(image_bytes: bytes) -> str:
    """将图像二进制流编码为 base64 字符串"""
    return base64.b64encode(image_bytes).decode('utf-8')


def get_image_mime_type(filename: str) -> str:
    """根据文件后缀获取 MIME 类型"""
    ext = os.path.splitext(filename)[1].lower()
    if ext in ['.jpg', '.jpeg']:
        return 'image/jpeg'
    elif ext == '.png':
        return 'image/png'
    elif ext == '.webp':
        return 'image/webp'
    elif ext == '.gif':
        return 'image/gif'
    return 'image/jpeg'


def parse_image_with_vlm(
    image_bytes: bytes,
    filename: str,
    api_key: str,
    base_url: Optional[str] = None,
    model_name: str = "gpt-4o",
    custom_prompt: Optional[str] = None
) -> str:
    """
    调用视觉大语言模型 (VLM) 对图像进行 OCR 与高阶多模态理解
    - 识别图像中的标题、段落文字
    - 将图中的表格转换为标准 Markdown 表格
    - 针对图表/架构图/流程图生成结构化的逻辑说明
    """
    from openai import OpenAI

    client = OpenAI(
        api_key=api_key,
        base_url=base_url if base_url else None
    )

    mime_type = get_image_mime_type(filename)
    b64_img = encode_image_to_base64(image_bytes)

    system_prompt = (
        "你是一名顶尖的高精度多模态文档与图像解析专家（OCR & VLM）。\n"
        "请将图像内容完整、准确地转换为结构化的 Markdown 文本：\n"
        "1. 保留所有印刷文字与手写体，严格保证字符准确性；\n"
        "2. 图中的数据表格必须转换为标准的 Markdown 表格；\n"
        "3. 数学公式使用 LaTeX 格式（如 $E=mc^2$）；\n"
        "4. 如果是流程图、架构图或统计图表，请在还原所有文字的同时，详细描述图表结构、数据走势和内在逻辑；\n"
        "5. 直接输出解析后的 Markdown 内容，不要带有额外的闲聊。"
    )

    user_prompt = custom_prompt or "请解析此图像中的所有文本、表格和图表内容，并输出为规范的 Markdown。"

    messages = [
        {"role": "system", "content": system_prompt},
        {
            "role": "user",
            "content": [
                {"type": "text", "text": user_prompt},
                {
                    "type": "image_url",
                    "image_url": {
                        "url": f"data:{mime_type};base64,{b64_img}"
                    }
                }
            ]
        }
    ]

    try:
        response = client.chat.completions.create(
            model=model_name,
            messages=messages,
            temperature=0.1
        )
        return response.choices[0].message.content.strip()
    except Exception as e:
        return f"[VLM 解析失败]: {str(e)}"


def parse_document_file(
    filename: str,
    file_bytes: bytes,
    vlm_api_key: Optional[str] = None,
    vlm_base_url: Optional[str] = None,
    vlm_model: str = "gpt-4o"
) -> List[Document]:
    """
    通用多格式文档解析器：
    支持 PDF, DOCX, XLSX/XLS, CSV, JSON, TXT, MD 及图片格式
    输出标准 LangChain Document 列表
    """
    ext = os.path.splitext(filename)[1].lower()
    docs: List[Document] = []
    file_stream = io.BytesIO(file_bytes)

    # 1. PDF 格式
    if ext == '.pdf':
        reader = pypdf.PdfReader(file_stream)
        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                docs.append(Document(
                    page_content=page_text.strip(),
                    metadata={
                        "source": filename,
                        "page": idx + 1,
                        "file_type": "pdf"
                    }
                ))

    # 2. Word 格式 (.docx)
    elif ext == '.docx':
        doc = docx.Document(file_stream)
        elements = []
        for p in doc.paragraphs:
            if p.text.strip():
                elements.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_str = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_str:
                    elements.append(row_str)
        content = "\n\n".join(elements)
        if content.strip():
            docs.append(Document(
                page_content=content,
                metadata={"source": filename, "file_type": "docx"}
            ))

    # 3. Excel 表格 (.xlsx, .xls)
    elif ext in ['.xlsx', '.xls']:
        excel_data = pd.read_excel(file_stream, sheet_name=None)
        for sheet_name, df in excel_data.items():
            if not df.empty:
                table_md = df.to_markdown(index=False)
                docs.append(Document(
                    page_content=f"### 工作表 (Sheet): {sheet_name}\n\n{table_md}",
                    metadata={"source": filename, "sheet": sheet_name, "file_type": "excel"}
                ))

    # 4. CSV 表格
    elif ext == '.csv':
        df = pd.read_csv(file_stream)
        if not df.empty:
            docs.append(Document(
                page_content=df.to_markdown(index=False),
                metadata={"source": filename, "file_type": "csv"}
            ))

    # 5. JSON 数据
    elif ext == '.json':
        try:
            raw_str = file_bytes.decode('utf-8', errors='ignore')
            data = json.loads(raw_str)
            content = json.dumps(data, ensure_ascii=False, indent=2)
        except Exception:
            content = file_bytes.decode('utf-8', errors='ignore')
        docs.append(Document(
            page_content=content,
            metadata={"source": filename, "file_type": "json"}
        ))

    # 6. TXT / Markdown 纯文本
    elif ext in ['.txt', '.md']:
        content = file_bytes.decode('utf-8', errors='ignore')
        if content.strip():
            docs.append(Document(
                page_content=content.strip(),
                metadata={"source": filename, "file_type": ext.replace('.', '')}
            ))

    # 7. 图片与扫描件 (调用 VLM 解析)
    elif ext in ['.png', '.jpg', '.jpeg', '.webp', '.bmp']:
        if vlm_api_key:
            parsed_md = parse_image_with_vlm(
                image_bytes=file_bytes,
                filename=filename,
                api_key=vlm_api_key,
                base_url=vlm_base_url,
                model_name=vlm_model
            )
            docs.append(Document(
                page_content=f"## [多模态图像/扫描件解析: {filename}]\n\n{parsed_md}",
                metadata={"source": filename, "file_type": "image_vlm"}
            ))
        else:
            docs.append(Document(
                page_content=f"[图像文件: {filename}。未配置 VLM API Key，未能提取图内文本]",
                metadata={"source": filename, "file_type": "image_placeholder"}
            ))
    else:
        print(f"[vlm_ocr] 警告: 暂不支持解析扩展名为 '{ext}' 的文件: {filename}")

    return docs
